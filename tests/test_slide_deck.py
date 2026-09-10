"""The `slide_deck` capability: a generation that leaves a file behind.

#35's argument, pinned: a deck is not a report rendered differently, so the
deck-shaped structure is asked of the model (here, on the capability side,
where the tokens are booked) and only then rendered. Everything below runs
without python-pptx and without a disk — that is what the `DeckRenderer` and
`ArtifactStore` ports are for; the one test that needs the real library says
so with `importorskip`.
"""
from __future__ import annotations

import json
from io import BytesIO
from pathlib import Path

import pytest
from conftest import FakeLLM, plan_json

from jobsmith.agents.base import AgentContext
from jobsmith.agents.default import DefaultResources, default_capabilities
from jobsmith.agents.default.slides import (
    MAX_BULLETS,
    MAX_CHARS,
    MAX_SLIDES,
    Deck,
    Slide,
    SlideDeckCapability,
)
from jobsmith.core.artifacts import LocalArtifactStore, artifact_refs
from jobsmith.core.builder import build_agent
from jobsmith.core.deps import Deps
from jobsmith.core.registry import CapabilityRegistry
from jobsmith.jobs.manager import JobManager
from jobsmith.jobs.models import JobStatus

DESIGNED = json.dumps({
    "title": "Background jobs",
    "subtitle": "what we learned",
    "slides": [
        {"title": "The problem", "bullets": ["long waits", "no feedback"],
         "notes": "Open on the pain."},
        {"title": "What we did", "bullets": ["a job engine"], "notes": ""},
    ],
})


class FakeRenderer:
    """`DeckRenderer` that keeps the deck instead of rendering it."""

    extension = "deck"

    def __init__(self, fail: str = ""):
        self.fail = fail
        self.seen: list[Deck] = []

    async def render(self, deck: Deck) -> bytes:
        if self.fail:
            raise RuntimeError(self.fail)
        self.seen.append(deck)
        return f"[{deck.title}]".encode()


class MemoryStore:
    """`ArtifactStore` with no disk: the port exists so a test needs none."""

    def __init__(self, fail: str = ""):
        self.fail = fail
        self.written: dict[tuple[str, str], bytes] = {}

    async def write(self, job_id: str, name: str, data: bytes | str) -> str:
        if self.fail:
            raise OSError(self.fail)
        payload = data.encode("utf-8") if isinstance(data, str) else data
        self.written[(job_id, name)] = payload
        return f"memory://{job_id}/{name}"


def make_capability(reply: str = DESIGNED, *, renderer=None, store=None):
    llm = FakeLLM({"Design a slide deck": reply}, default=reply)
    cap = SlideDeckCapability(llm, store or MemoryStore(), renderer or FakeRenderer())
    return cap, llm


async def run_capability(cap, **state):
    return await cap.build().ainvoke({"query": "brief the board", "results": {}, **state})


def result_of(out) -> dict:
    return out["results"]["slide_deck"]


# ---------------------------------------------------------------- the design

async def test_the_model_is_asked_for_a_deck_and_the_file_is_declared():
    """The whole capability in one run: design → render → a declared file."""
    renderer, store = FakeRenderer(), MemoryStore()
    cap, llm = make_capability(renderer=renderer, store=store)
    result = result_of(await run_capability(cap, job_id="job1"))

    assert result["ok"] is True
    assert result["data"]["path"] == "memory://job1/slide_deck.deck"
    assert store.written[("job1", "slide_deck.deck")] == b"[Background jobs]"
    # the renderer was handed a structure, never prose or markup
    (deck,) = renderer.seen
    assert deck.title == "Background jobs" and deck.subtitle == "what we learned"
    assert [s.title for s in deck.slides] == ["The problem", "What we did"]
    assert deck.slides[0].bullets == ("long waits", "no feedback")
    assert deck.slides[0].notes == "Open on the pain."
    # declared as an artifact, titled for a human, and NOT via the fallback
    assert artifact_refs(result["meta"])[0].path == result["data"]["path"]
    assert artifact_refs(result["meta"])[0].title == "Background jobs"
    assert result["meta"]["via_fallback"] is False and result["meta"]["slide_count"] == 2
    # JSON was asked for, and the material went with the request
    assert llm.calls[0]["response_format"] == {"type": "json_object"}


async def test_the_deck_is_built_from_what_the_other_steps_produced():
    """A deck of the job's material, in a fixed order — `results` arrives in
    wave order and a consumer must never iterate that (core/state.py)."""
    cap, llm = make_capability()
    await run_capability(cap, job_id="job1", results={
        "critique": {"ok": True, "data": {"critique": "the weak point"}},
        "research": {"ok": True, "data": {"notes": "the notes"}},
        "analysis": {"ok": True, "data": {"analysis": "the finding"}},
        "documents": {"ok": False, "error": "nothing found"},
    })
    user = llm.calls[0]["messages"][-1]["content"]
    assert user.index("[analysis") < user.index("[research") < user.index("[critique")
    assert "the finding" in user and "brief the board" in user
    assert "nothing found" not in user            # a failed step has no material


async def test_the_internal_review_reaches_the_deck_labelled_as_a_review():
    """#58: the deck is the one deliverable that reads the material directly.

    `critique` is agent-facing by design — its own prompt asks it to challenge
    the work and suggest improvements — so a block tagged only `[critique]` is
    material the model has no reason to treat differently, and it did not: the
    run that opened #58 shipped slides titled *État actuel et risques* and
    *Prochaines étapes* to someone who had asked about a subject. The block now
    carries what it *is*, and `DESIGN_SYSTEM` says what to do with a block of
    that kind. Read off `MATERIAL` rather than retyped here: the label and the
    prompt are the fix, and a test that copies them proves neither travelled.
    """
    roles = {name: role for name, _key, role in SlideDeckCapability.MATERIAL}
    assert "OF THE WORK" in roles["critique"]

    cap, llm = make_capability()
    await run_capability(cap, job_id="job1", results={
        "critique": {"ok": True, "data": {"critique": "the weak point"}},
    })
    system, user = (m["content"] for m in llm.calls[0]["messages"])
    assert f"[critique — {roles['critique']}]" in user
    assert "the weak point" in user
    assert "labelled with what it is" in system     # the rule for such a block


async def test_a_deck_with_no_material_still_gets_the_request():
    cap, llm = make_capability()
    await run_capability(cap, job_id="job1")
    assert "no upstream material" in llm.calls[0]["messages"][-1]["content"]


def test_the_structure_is_read_back_leniently_and_bounded():
    """It comes straight from a model, so every field is a suggestion."""
    deck = Deck.from_dict({
        "title": "T" * (MAX_CHARS + 50),
        "slides": [
            {"title": "kept", "bullets": ["a"] * (MAX_BULLETS + 4), "notes": "n"},
            {"bullets": ["no title is fine"]},
            {"title": "", "bullets": []},          # nothing on it — not a slide
            "not a slide at all",
            *[{"title": f"s{i}"} for i in range(MAX_SLIDES + 5)],
        ],
    })
    assert len(deck.title) == MAX_CHARS
    assert len(deck.slides) == MAX_SLIDES
    assert len(deck.slides[0].bullets) == MAX_BULLETS
    assert deck.slides[1].title == "…"             # a bullet list needs a heading
    assert Deck.from_dict("nonsense").slides == ()
    assert Deck.from_dict({"slides": {"not": "a list"}}).slides == ()


# ---------------------------------------------------------------- degradation

async def test_a_reply_that_is_not_json_still_leaves_a_deck():
    """Lenient like `documents` and `research` — and it says it degraded."""
    renderer = FakeRenderer()
    cap, _ = make_capability("# The finding\n- it is slow\n- and opaque\n\nSo: fix it",
                             renderer=renderer)
    result = result_of(await run_capability(cap, job_id="job1"))

    assert result["ok"] is True and result["meta"]["via_fallback"] is True
    (deck,) = renderer.seen
    assert [s.title for s in deck.slides] == ["The finding", "So: fix it"]
    assert deck.slides[0].bullets == ("it is slow", "and opaque")


async def test_json_wrapped_in_a_fence_is_still_a_designed_deck():
    cap, _ = make_capability(f"here you go:\n```json\n{DESIGNED}\n```")
    result = result_of(await run_capability(cap, job_id="job1"))
    assert result["meta"]["via_fallback"] is False
    assert result["data"]["deck"]["title"] == "Background jobs"


async def test_nothing_to_put_on_a_slide_is_a_failed_step():
    store = FakeRenderer(), MemoryStore()
    cap, _ = make_capability("   ", renderer=store[0], store=store[1])
    out = await run_capability(cap, job_id="job1")
    result = result_of(out)

    assert result["ok"] is False and "nothing to put on a slide" in result["error"]
    assert store[1].written == {} and store[0].seen == []
    assert artifact_refs(result["meta"]) == []     # nothing was written to declare
    assert out["errors"][0]["recoverable"] is True  # the run goes on without it


# ---------------------------------------------------------------- no job, no file

async def test_a_run_outside_a_job_writes_nothing_and_spends_nothing():
    """`job_id` is NOT Required on CapabilityBaseState: a graph driven outside
    a job has none, and a deck is a file. Refused BEFORE the model call — a
    deck nobody could ever open is not worth designing."""
    renderer, store = FakeRenderer(), MemoryStore()
    cap, llm = make_capability(renderer=renderer, store=store)
    result = result_of(await run_capability(cap))          # no job_id in state

    assert result["ok"] is False and "no job to write a deck for" in result["error"]
    assert llm.calls == [] and store.written == {} and renderer.seen == []


# ---------------------------------------------------------------- the write

async def test_a_deck_that_could_not_be_written_declares_no_file():
    """No path came back, so there is nothing to declare: a ref to a file that
    may or may not exist is what the manager refuses to record (#40)."""
    cap, _ = make_capability(store=MemoryStore(fail="No space left on device"))
    result = result_of(await run_capability(cap, job_id="job1"))

    assert result["ok"] is False
    assert "could not be written" in result["error"]
    assert "No space left on device" in result["error"]
    assert artifact_refs(result["meta"]) == []


async def test_a_renderer_that_blows_up_fails_the_step_and_not_the_run():
    cap, _ = make_capability(renderer=FakeRenderer(fail="corrupt template"))
    out = await run_capability(cap, job_id="job1")
    assert result_of(out)["ok"] is False
    assert "corrupt template" in result_of(out)["error"]
    assert out["errors"][0]["source"] == "slide_deck"


async def test_a_file_already_on_disk_is_declared_even_when_the_step_fails():
    """`_emit_failure` takes the same `meta` as `_emit_success` (#41).

    Nothing between the store's write and the emit can fail today, so this is
    a guarantee about tomorrow: the day a node lands after `render`, the deck
    it wrote must still be declared rather than orphaned.
    """
    cap, _ = make_capability()
    emitted = await cap.emit_failure(
        {"query": "q", "failure": "died after the write", "path": "memory://job1/d.deck"})
    result = emitted["results"]["slide_deck"]

    assert result["ok"] is False and result["error"] == "died after the write"
    assert [r.path for r in artifact_refs(result["meta"])] == ["memory://job1/d.deck"]


# ---------------------------------------------------------------- presentation

async def test_the_capability_presents_its_own_deck():
    cap, _ = make_capability()
    result = result_of(await run_capability(cap, job_id="job1"))

    report = cap.render_report(result)
    assert "**Background jobs** — 2 slides" in report
    assert "`memory://job1/slide_deck.deck`" in report      # the file, for a human
    assert "1. The problem" in report and "2. What we did" in report
    assert "long waits" not in report                       # an outline, not the deck

    # ...and the generator gets the fact, not the outline: a written answer
    # handed the slide titles transcribes them instead of referring to them,
    # and the deck's register comes with them (#58)
    context = cap.render_context(result)
    assert "Slide deck produced" in context and "2 slides" in context
    assert "The problem" not in context and "long waits" not in context
    assert cap.render_report({"ok": False, "error": "boom"}) == "_boom_"
    assert cap.render_context({"ok": True, "data": {}}) is None


# ---------------------------------------------------------------- registration

def make_context(*, renderer=None, artifacts=None) -> AgentContext:
    return AgentContext(llm=FakeLLM(), resources=DefaultResources(decks=renderer),
                        artifacts=artifacts)


def test_the_deck_step_stays_out_of_the_registry_when_nothing_can_serve_it():
    """The pack's standing rule: no renderer (no `.[pptx]`) or nowhere to put
    the file means no capability, rather than a step the planner can only
    watch fail."""
    def names(ctx) -> list[str]:
        return [c.spec.name for c in default_capabilities(ctx)]

    assert "slide_deck" not in names(make_context(artifacts=MemoryStore()))
    assert "slide_deck" not in names(make_context(renderer=FakeRenderer()))
    both = names(make_context(renderer=FakeRenderer(), artifacts=MemoryStore()))
    assert both[-1] == "slide_deck"      # last: it presents what the others made
    assert both[:3] == ["research", "analysis", "critique"]


# ---------------------------------------------------------------- the renderer

def test_the_pptx_renderer_writes_a_real_presentation():
    """The adapter, against the library it adapts."""
    pytest.importorskip("pptx")
    from pptx import Presentation

    from jobsmith.agents.default.pptx_deck import PptxRenderer

    data = PptxRenderer().build(Deck(
        title="Background jobs", subtitle="what we learned",
        slides=(Slide("The problem", ("long waits", "no feedback"), "Open on the pain."),
                Slide("No bullets here", ()))))

    assert data[:2] == b"PK" and PptxRenderer.extension == "pptx"
    slides = list(Presentation(BytesIO(data)).slides)
    assert [[s.text_frame.text for s in slide.shapes if s.has_text_frame]
            for slide in slides] == [
        ["Background jobs", "what we learned"],
        ["The problem", "long waits\nno feedback"],
        ["No bullets here"],          # an empty placeholder would prompt the reader
    ]
    # speaker notes are the half of a deck the slide never shows
    assert slides[1].notes_slide.notes_text_frame.text == "Open on the pain."
    assert not slides[2].has_notes_slide


def test_a_deck_is_16_9_and_its_content_moved_with_it():
    """#62: the stock template is 4:3, and resizing it is not a one-liner.

    Every deck jobsmith produced was 10 × 7.5 in — python-pptx's default, and
    the default nowhere else for fifteen years. Setting `slide_width` alone
    leaves the template's placeholders where a 10-inch canvas put them, so the
    property worth pinning is not "the slide is wide" but "the content is
    still centred in it": measured on the naive version, a body ended 3.8
    inches short of the right edge on every slide.
    """
    pytest.importorskip("pptx")
    from pptx import Presentation

    from jobsmith.agents.default.pptx_deck import PptxRenderer

    data = PptxRenderer().build(Deck(
        title="Background jobs", subtitle="what we learned",
        slides=(Slide("The problem", ("long waits", "no feedback")), Slide("Bare"))))

    deck = Presentation(BytesIO(data))
    assert deck.slide_width / deck.slide_height == pytest.approx(16 / 9, abs=0.001)
    assert (deck.slide_width, deck.slide_height) == (12192000, 6858000)
    for slide in deck.slides:
        for placeholder in slide.placeholders:
            right = deck.slide_width - (placeholder.left + placeholder.width)
            assert placeholder.left == right, "the content did not follow the page"


# ---------------------------------------------------------------- in a job

async def test_the_job_records_the_deck_as_an_annex(store, checkpointer, tmp_path):
    """End to end: the file a step produced is one of the job's outputs, after
    the deliverable and attributed to the step (#35 step 1's seam)."""
    llm = FakeLLM({"planner": plan_json("slide_deck"), "Design a slide deck": DESIGNED},
                  default="A sufficiently long final answer for the deck test.")
    cap = SlideDeckCapability(llm, LocalArtifactStore(tmp_path / "artifacts"), FakeRenderer())
    graph = build_agent(Deps(llm=llm), CapabilityRegistry([cap]), checkpointer=checkpointer)
    manager = JobManager(graph, store, reports_dir=tmp_path / "artifacts")

    done = await manager.run_job((await manager.create_job("brief the board")).job_id)

    assert done.status is JobStatus.DONE
    main, annex = done.outputs
    assert (main.role, main.format) == ("main", "markdown")
    assert (annex.role, annex.format, annex.produced_by) == ("annex", "deck", "slide_deck")
    assert annex.title == "Background jobs"
    assert Path(annex.path).read_bytes() == b"[Background jobs]"
    assert annex.path == str(tmp_path / "artifacts" / done.job_id / "slide_deck.deck")
    assert done.error is None                     # the file it promised is there
    assert done.report_path == main.path           # the deck did not become the report
